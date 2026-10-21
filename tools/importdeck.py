#!/usr/bin/env python3
"""Bring a PowerPoint or a PDF into a deck made with Import a deck (New deck › Import a deck, or
`./sb new "<name>" --from <file>`). The file is in the deck's source/ folder.

    python3 tools/importdeck.py --check            the Python packages this needs, and LibreOffice
    python3 tools/importdeck.py <deck>             keep each slide's look: every slide comes into the
                                                   deck file with its title and speaker notes, and every
                                                   word on it is editable
    python3 tools/importdeck.py <deck> --outline   for a rebuild: source/outline.md, each slide's title,
                                                   text, tables, chart numbers, and notes, with a picture
                                                   of each slide in source/pages/

<deck> is the deck's slug (./sb decks lists them) or its folder.

A PowerPoint file is converted by tools/pptx2native.py: text boxes, boxes, lines, tables, and pictures
become native shapes; charts, SmartArt, and icons are cut from a PDF of the deck as pictures. A PDF
on its own keeps each page's look too: the page is drawn without its words, and the words go back
on top as text you can edit, line by line.

Options:
    --pdf FILE       the deck exported as PDF, to render from. Default: a PDF in source/, or one
                     LibreOffice makes when it is installed.
    --no-pdf         convert a PowerPoint file without a PDF: charts, SmartArt, and icons are left
                     out, and there are no pictures of the original slides to compare with
    --slides 3-9,12  only these slides
    --no-wire        convert into ct-slides.js and img/ only, and leave the deck file's slides alone
                     (for a rebuild that keeps a few slides as they were: use ct('s7', ...) for those)
    --force          write the slides into the deck file even though they changed since the import

Needs python-pptx, lxml, Pillow, and PyMuPDF:
    python3 -m pip install python-pptx lxml Pillow PyMuPDF
"""
import argparse
import html
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path

APP = Path(__file__).resolve().parent.parent
sys.path[:0] = [str(APP), str(APP / 'tools')]
from newdeck import js_str, slugify   # noqa: E402
try:
    import sbconfig   # noqa: E402
    REGISTRY, tilde = sbconfig.REGISTRY, sbconfig.tilde
except ImportError:   # an app without sbconfig.py keeps decks.json next to it
    REGISTRY = APP / 'decks.json'

    def tilde(path):
        path = Path(path).resolve()
        return '~/' + str(path.relative_to(Path.home())) if Path.home() in path.parents else str(path)

PREFIX = 'imp'   # img/imp-001.jpg is slide 1 as it looked; img/imp-pic-*.png are the pictures on it
START = '// ---- IMPORTED SLIDES'
END = '// ---- END OF IMPORTED SLIDES ----'
PACKAGES = (('pptx', 'python-pptx'), ('lxml', 'lxml'), ('PIL', 'Pillow'), ('fitz', 'PyMuPDF'))
INSTALL = 'python3 -m pip install python-pptx lxml Pillow PyMuPDF'
P14 = '{http://schemas.microsoft.com/office/powerpoint/2010/main}'


def missing():
    out = []
    for mod, pkg in PACKAGES:
        try:
            __import__(mod)
        except ImportError:
            out.append(pkg)
    return out


def soffice():
    for name in ('soffice', 'libreoffice'):
        if shutil.which(name):
            return shutil.which(name)
    mac = Path('/Applications/LibreOffice.app/Contents/MacOS/soffice')
    return str(mac) if mac.exists() else None


def find_deck(arg):
    reg = json.loads(REGISTRY.read_text(encoding='utf-8')) if REGISTRY.exists() else {'decks': []}
    p = Path(os.path.expanduser(arg))
    for d in reg['decks']:
        folder = Path(os.path.expanduser(d['folder'])).resolve()
        if d['slug'] == arg or (p.is_dir() and folder == p.resolve()):
            return d, folder
    if p.is_dir() and (p / 'deck.html').exists():
        return {'slug': p.name, 'entry': 'deck.html'}, p.resolve()
    sys.exit(f'No deck {arg!r}. Give its slug (./sb decks lists them) or its folder.')


def source_of(deck, folder):
    if deck.get('source') and (folder / deck['source']).is_file():
        return folder / deck['source']
    files = sorted(p for p in (folder / 'source').glob('*') if p.suffix.lower() in ('.pptx', '.pdf'))
    return next((p for p in files if p.suffix.lower() == '.pptx'), files[0] if files else None)


NO_PDF = """There's no PDF of this deck in {src}, and LibreOffice isn't installed to make one.
The converter cuts charts, SmartArt, and icons out of a PDF of each slide, and the pictures of the
original slides come from it. Any one of these works:
  - export a PDF: in PowerPoint, File › Export › PDF (Keynote: File › Export To › PDF; Google Slides:
    File › Download › PDF), and save it in {src}
  - install LibreOffice (libreoffice.org) and run this again: it makes the PDF itself
  - run again with --no-pdf: every word, box, line, table, and picture comes in, but charts,
    SmartArt, and icons are left out"""


def pdf_for(src, args, say, required=True):
    if args.pdf:
        p = Path(os.path.expanduser(args.pdf))
        return p if p.is_file() else sys.exit(f'No PDF at {args.pdf}')
    if src.suffix.lower() == '.pdf':
        return src
    if args.no_pdf:
        return None
    for p in [src.with_suffix('.pdf')] + sorted(src.parent.glob('*.pdf')) + sorted(src.parent.glob('*.PDF')):
        if p.is_file():
            return p
    so = soffice()
    if so:
        say(f'Making a PDF with LibreOffice ({so}). A big deck takes a minute or two.')
        try:
            subprocess.run([so, '--headless', '--convert-to', 'pdf', '--outdir', str(src.parent), str(src)], capture_output=True, timeout=1200)
        except (OSError, subprocess.SubprocessError) as e:
            say(f'LibreOffice could not make it: {e}')
        if src.with_suffix('.pdf').is_file():
            return src.with_suffix('.pdf')
    if not required:   # the outline still has every word; only the pictures of the slides need a PDF
        say(f'No PDF of the deck in {tilde(src.parent)} (and no LibreOffice to make one), so the outline has no pictures of the slides.')
        return None
    sys.exit(NO_PDF.format(src=tilde(src.parent)))


def page_map(slides, hidden, pages, wanted):
    """Which PDF page shows each slide. PowerPoint and LibreOffice leave hidden slides out of a PDF."""
    if pages == slides:
        return {n: n - 1 for n in wanted}
    shown = [n for n in range(1, slides + 1) if n not in hidden]
    if pages == len(shown):
        at = {n: i for i, n in enumerate(shown)}
        return {n: at[n] for n in wanted if n in at}
    sys.exit(f'The PDF has {pages} pages and the deck {slides} slides ({len(hidden)} hidden), so its pages can\'t be matched '
             'to the slides. Export the PDF again from this same file.')


def clean(text, limit=None):
    t = re.sub(r'\s+', ' ', (text or '').replace('\x0b', ' ')).strip()
    return t[:limit].rsplit(' ', 1)[0] + '…' if limit and len(t) > limit else t


def notes_html(text):
    paras = [p.strip() for p in re.split(r'\n\s*', (text or '').replace('\x0b', '\n')) if p.strip()]
    return html.escape('\n'.join(paras), quote=False)


def title_shape(shapes):
    """No title placeholder: the text box near the top of the converted slide with the biggest words, and its first paragraph."""
    best = None
    for sh in shapes or []:
        if sh.get('k') != 'box' or not sh.get('text') or sh.get('y', 999) > 300:
            continue
        first = next((''.join(r.get('t', '') for r in p.get('runs') or []).strip() for p in sh['text']
                      if ''.join(r.get('t', '') for r in p.get('runs') or []).strip()), '')
        size = max([r.get('sz') or 0 for p in sh['text'][:1] for r in p.get('runs') or []] + [0])
        if first and (best is None or (size, -sh['y']) > best[0]):
            best = ((size, -sh['y']), sh, first)
    return (best[1], clean(best[2], 120)) if best else (None, '')


def title_from_shapes(shapes):
    return title_shape(shapes)[1]


def slide_id(n, title):
    """s7-the-first-few-words-of-its-title: unique by its number, readable by its words."""
    words, out = slugify(title).split('-') if title else [], f's{n}'
    for w in words:
        if len(out) + len(w) + 1 > 34:
            break
        out += '-' + w
    return out


# ---------- PowerPoint ----------

def title_from_pptx(slide, height):
    """No title placeholder and nothing converted to look at: the first line of the text near the top set biggest."""
    best = None
    for sh in slide.shapes:
        if not getattr(sh, 'has_text_frame', False) or not sh.has_text_frame or (sh.top or 0) > (height or 6858000) * 0.4:
            continue
        para = next((p for p in sh.text_frame.paragraphs if p.text.strip()), None)
        if para is None:
            continue
        size = max([r.font.size.pt for r in para.runs if r.font.size] + [0])
        if best is None or (size, -(sh.top or 0)) > best[0]:
            best = ((size, -(sh.top or 0)), clean(para.text, 120))
    return best[1] if best else ''


def pptx_facts(prs, data):
    """Each slide's title, notes, section, and whether it is hidden."""
    sections = {}
    for sec in prs.part._element.iter(P14 + 'section'):
        for sid in sec.iter(P14 + 'sldId'):
            sections[int(sid.get('id'))] = sec.get('name') or ''
    out = {}
    for n, s in enumerate(prs.slides, 1):
        title = ''
        try:
            t = s.shapes.title
            title = clean(t.text_frame.text, 140) if t is not None and t.has_text_frame else ''
        except (AttributeError, KeyError, ValueError):
            pass
        if not title:
            title = title_from_shapes((data.get(n) or {}).get('shapes')) or title_from_pptx(s, prs.slide_height)
        notes = ''
        if s.has_notes_slide and s.notes_slide.notes_text_frame is not None:
            notes = s.notes_slide.notes_text_frame.text
        out[n] = {'title': title, 'notes': notes.strip(), 'section': sections.get(s.slide_id),
                  'hidden': s._element.get('show') == '0', 'layout': s.slide_layout.name}
    return out


def convert_pptx(src, folder, pdf, wanted, say):
    import pptx2native as P
    img = folder / 'img'
    img.mkdir(exist_ok=True)
    tmp = tempfile.mkdtemp(prefix='importdeck-')
    try:
        conv = P.Converter(str(src), str(img), PREFIX, str(img), tmp)
        conv.keep_fonts = True   # the deck's own typefaces, where this machine has them
        n = len(conv.prs.slides)
        wanted = {k for k in wanted if 1 <= k <= n} if wanted else set(range(1, n + 1))
        pics = {}
        if pdf:
            import fitz
            hidden = {i for i, s in enumerate(conv.prs.slides, 1) if s._element.get('show') == '0'}
            pages = page_map(n, hidden, fitz.open(str(pdf)).page_count, wanted)
            say(f'Rendering {len(pages)} pages of {pdf.name}…')
            P.render_pdf(str(pdf), str(img), tmp, PREFIX, pages, (conv.prs.slide_width, conv.prs.slide_height))
            pics = {k: f'img/{PREFIX}-{k:03d}.jpg' for k in pages}
            conv.pdf_fonts = {re.sub(r'[^a-z0-9]', '', re.sub(r'^[A-Z]{6}\+', '', f[3]).lower()) for page in fitz.open(str(pdf)) for f in page.get_fonts()}
        say(f'Converting {len(wanted)} slides…')
        conv.only = wanted
        data = conv.run()
        for k, v in data.items():   # no title placeholder: mark the box the slide is named after, where Checks looks for a title
            shapes = v.get('shapes') or []
            if not any(sh.get('ph') == 'title' for sh in shapes) and not (conv.prs.slides[k - 1].shapes.title is not None):
                box, _ = title_shape(shapes)
                if box:
                    box['ph'] = 'title'
        return conv, data, pics, pptx_facts(conv.prs, data)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# ---------- PDF ----------

def pdf_size(page):
    """The PDF page's size as slide EMU, for fitting it onto the 16:9 canvas."""
    return round(page.rect.width * 12700), round(page.rect.height * 12700)


def pdf_font(name):
    """A PDF's font name as a CSS font family, the deck's own font after it: ABCDEF+Georgia-Bold -> 'Georgia'."""
    face = re.sub(r'^[A-Z]{6}\+', '', name or '').split('-')[0].split(',')[0]
    face = re.sub(r'(PSMT|MT|PS)$', '', face)
    face = re.sub(r'(?<=[a-z])(?=[A-Z])', ' ', face).strip()
    return f"'{face}', var(--slide)" if re.fullmatch(r'[\w .&]{2,60}', face) else None


def hexc(c):
    return '#%06X' % (c & 0xFFFFFF)


def convert_pdf(pdf, folder, wanted, say):
    """Each page: a picture of the page without its words, and the words on top as native text,
    one line per text box, at the page's own sizes, weights, and colors."""
    import fitz
    from PIL import Image
    import pptx2native as P
    img = folder / 'img'
    img.mkdir(exist_ok=True)
    doc = fitz.open(str(pdf))
    n = doc.page_count
    wanted = {k for k in wanted if 1 <= k <= n} if wanted else set(range(1, n + 1))
    tmp = tempfile.mkdtemp(prefix='importdeck-')
    data, pics, facts = {}, {}, {}
    try:
        size = pdf_size(doc[0])
        say(f'Rendering {len(wanted)} pages of {pdf.name}…')
        P.render_pdf(str(pdf), str(img), tmp, PREFIX, {k: k - 1 for k in wanted}, size)
        pics = {k: f'img/{PREFIX}-{k:03d}.jpg' for k in wanted}
        bare = fitz.open(str(pdf))   # the same pages with their words taken out
        for k in sorted(wanted):
            page, bpage = doc[k - 1], bare[k - 1]
            s, (dx, dy) = P.fit_slide(*pdf_size(page))
            k_px = 12700 * s / P.EMU_PER_PX                 # PDF points -> canvas px
            to_pt = s                                       # PDF points -> points on the 960pt-wide canvas
            ox, oy = P.emu_px(dx), P.emu_px(dy)
            shapes = []
            for b in page.get_text('dict')['blocks']:
                if b.get('type') != 0:
                    continue
                lines = [l for l in b.get('lines', []) if tuple(round(v, 2) for v in l['dir']) == (1.0, 0.0) and ''.join(sp['text'] for sp in l['spans']).strip()]
                if not lines:
                    continue
                x0s, x1s = [l['bbox'][0] for l in lines], [l['bbox'][2] for l in lines]
                align = 'l'
                if len(lines) > 1 and max(x0s) - min(x0s) > 3:
                    mids = [(a + z) / 2 for a, z in zip(x0s, x1s)]
                    align = 'r' if max(x1s) - min(x1s) <= 3 else ('ctr' if max(mids) - min(mids) <= 4 else 'l')
                for l in lines:
                    runs = []
                    for sp in l['spans']:
                        if not sp['text']:
                            continue
                        font = sp.get('font', '')
                        runs.append({'t': sp['text'], 'sz': round(sp['size'] * to_pt, 1), 'b': bool(sp['flags'] & 16) or 'bold' in font.lower() or 'black' in font.lower(),
                                     'i': bool(sp['flags'] & 2) or 'italic' in font.lower(), 'u': False, 'c': hexc(sp['color'])})
                        if pdf_font(font):
                            runs[-1]['font'] = pdf_font(font)
                        bpage.add_redact_annot(fitz.Rect(sp['bbox']), fill=False)
                    if not runs:
                        continue
                    x0, y0, x1, y1 = l['bbox']
                    pad = 6
                    shapes.append({'k': 'box', 'x': round(ox + x0 * k_px - (pad if align != 'l' else 0), 2), 'y': round(oy + y0 * k_px, 2),
                                   'w': round((x1 - x0) * k_px + (pad * 2 if align == 'ctr' else pad), 2), 'h': round((y1 - y0) * k_px, 2),
                                   'fill': None, 'line': None, 'rot': 0,
                                   'text': [{'runs': runs, 'align': align, 'bullet': '', 'level': 0, 'lnSpc': None, 'spcBef': None, 'spcAft': None,
                                             'marL': 0, 'indent': 0, 'buFont': None, 'buColor': None}],
                                   'anchor': 't', 'autofit': False, 'wrap': False, 'ins': [0, 0, 0, 0]})
            bpage.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE, graphics=fitz.PDF_REDACT_LINE_ART_NONE, text=fitz.PDF_REDACT_TEXT_REMOVE)
            W = 1600
            z = (W - 2 * ox * W / 1000) / page.rect.width
            pix = bpage.get_pixmap(matrix=fitz.Matrix(z, z), alpha=False)
            im = Image.frombytes('RGB', (pix.width, pix.height), pix.samples)
            if ox or oy:
                full = Image.new('RGB', (W, round(W * 0.5625)), im.getpixel((2, 2)))
                full.paste(im, (round(ox * W / 1000), round(oy * W / 1000)))
                im = full
            bg = f'img/{PREFIX}-bg-{k:03d}.jpg'
            im.save(folder / bg, quality=85, optimize=True, progressive=True)
            corner = '#%02X%02X%02X' % im.getpixel((2, 2))
            shapes.insert(0, {'k': 'pic', 'src': bg, 'x': 0, 'y': 0, 'w': 1000, 'h': 562.5, 'crop': [0, 0, 0, 0], 'radius': None, 'rot': 0})
            data[k] = {'bg': corner, 'bgpic': False, 'shapes': shapes}
            box, title = title_shape(shapes)
            if box:
                box['ph'] = 'title'   # where the title is, as a PowerPoint title placeholder would say (Checks looks for it)
            facts[k] = {'title': title, 'notes': '', 'section': None, 'hidden': False, 'layout': ''}
        return data, pics, facts
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# ---------- writing it into the deck ----------

def write_ct(folder, data, partial):
    """ct-slides.js: const CT = { s1: {...}, ... }. With --slides, the other slides stay."""
    path = folder / 'ct-slides.js'
    ct = {f's{k}': v for k, v in sorted(data.items())}
    if partial and path.exists():
        m = re.match(r'\s*const CT = (\{.*\});\s*$', path.read_text(encoding='utf-8'), re.S)
        try:
            old = json.loads(m.group(1)) if m else None
        except json.JSONDecodeError:
            old = None
        if old is None:   # edited by hand into something that isn't JSON: add to it instead
            with open(path, 'a', encoding='utf-8') as f:
                f.write('Object.assign(CT, ' + json.dumps(ct, ensure_ascii=False, separators=(',', ':')) + ');\n')
            return
        ct = dict(old, **ct)
    path.write_text('const CT = ' + json.dumps(ct, ensure_ascii=False, separators=(',', ':')) + ';\n', encoding='utf-8')


def library(src, data, pics, facts):
    """The CATS, CAT_LABELS, CAT_SHORT, and LIBRARY lines for the deck file."""
    order = sorted(data)
    cats, labels = [], {}
    for k in order:
        label = (facts[k].get('section') or '').strip()
        if label and label not in labels:
            cid = re.sub(r'-', '', slugify(label))[:20] or 'part'
            while cid in labels.values():
                cid += 'x'
            labels[label] = cid
            cats.append((cid, label))
    if not cats or any(not (facts[k].get('section') or '').strip() for k in order):
        cats.insert(0, ('deck', 'The deck'))
    cat_of = lambda k: labels.get((facts[k].get('section') or '').strip(), 'deck')
    lines = [f'{START}: tools/importdeck.py replaces everything from this line to the END line ----',
             f'// From source/{src.name}, {time.strftime("%B %-d, %Y")}. Each slide is ct(key, part, title, ...): its shapes are CT[key] in',
             '// ct-slides.js, src is a picture of the original slide (the "original" button), and notes are its speaker notes.',
             'const CATS = ' + json.dumps([c for c, _ in cats]) + ';',
             'const CAT_LABELS = { ' + ', '.join(f'{js_str(c)}: {js_str(lab)}' for c, lab in cats) + ' };',
             'const CAT_SHORT = { ' + ', '.join(f'{js_str(c)}: {js_str(clean(lab, 16))}' for c, lab in cats) + ' };',
             'const LIBRARY = {']
    used = set()
    for k in order:
        f = facts[k]
        title = f['title'] or f'Slide {k}'
        sid = slide_id(k, f['title'])
        while sid in used:
            sid += '-x'
        used.add(sid)
        extra = []
        if pics.get(k):
            extra.append(f'src:{js_str(pics[k])}')
        if f['hidden']:
            extra.append('alt:true')
        extra.append(f'notes:{js_str(notes_html(f["notes"]))}')
        lines.append(f"  '{sid}': ct('s{k}', {js_str(cat_of(k))}, {js_str(html.escape(title, quote=False))}, {{ {', '.join(extra)} }}),")
    lines += ['};', END]
    return '\n'.join(lines) + '\n'


def imported_block(f, force):
    """Where the imported slides go in the deck file: (its text, start, end). Checked before converting anything."""
    t = f.read_text(encoding='utf-8')
    i, j = t.find(START), t.find(END)
    if i < 0 or j < 0:
        sys.exit(f'{f.name} has no IMPORTED SLIDES lines, so it was not made by Import a deck. Run with --no-wire to convert into '
                 'ct-slides.js and img/, then write LIBRARY entries for the slides by hand (see ct() in templates/_import/content.js).')
    i = t.rfind('\n', 0, i) + 1
    j = t.find('\n', j) + 1 or len(t)
    if "'import-start'" not in t[i:j] and not force:
        sys.exit(f'The slides in {f.name} have changed since the import. To bring them in again over those changes (git keeps the '
                 'old ones), run this again with --force. To convert without touching the deck file, use --no-wire.')
    return t, i, j


HIDE_ORIGINAL = '.ctrls .cbtn.orig { display: none; }   /* no pictures of the original slides (imported without a PDF) */'


def wire(f, block, force, originals=True):
    t, i, j = imported_block(f, force)
    t = t[:i] + block + t[j:]
    # each slide's "original" button shows its picture: without pictures, it goes
    t = t.replace('\n' + HIDE_ORIGINAL, '')
    if not originals:
        t = t.replace('\n<style id="deck-theme">', '\n<style id="deck-theme">\n' + HIDE_ORIGINAL, 1)   # the tag itself, not a comment naming it
    f.write_text(t, encoding='utf-8')


# ---------- the outline, for a rebuild ----------

def shape_text(sh, depth=0):
    """A shape's words, tables, and chart numbers as Markdown lines (groups opened up)."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    out = []
    kind = getattr(sh, 'shape_type', None)
    if kind == MSO_SHAPE_TYPE.GROUP:
        for c in sorted(sh.shapes, key=lambda c: (c.top or 0, c.left or 0)):
            out += shape_text(c, depth)
        return out
    if getattr(sh, 'has_text_frame', False) and sh.has_text_frame:
        for p in sh.text_frame.paragraphs:
            t = clean(p.text)
            if t:
                out.append('  ' * p.level + '- ' + t)
    if getattr(sh, 'has_table', False) and sh.has_table:
        rows = [[clean(c.text).replace('|', '/') for c in r.cells] for r in sh.table.rows]
        if rows:
            out += ['', '| ' + ' | '.join(rows[0]) + ' |', '|' + '---|' * len(rows[0])] + ['| ' + ' | '.join(r) + ' |' for r in rows[1:]] + ['']
    if getattr(sh, 'has_chart', False) and sh.has_chart:
        ch = sh.chart
        name = ''
        try:
            name = clean(ch.chart_title.text_frame.text) if ch.has_title else ''
        except (AttributeError, ValueError):
            pass
        out.append(f'- Chart ({str(ch.chart_type).split(".")[-1].split(" ")[0].lower()}){": " + name if name else ""}')
        try:
            cats = [str(c) for c in ch.plots[0].categories]
            out.append('  - categories: ' + ', '.join(cats))
            for plot in ch.plots:
                for s in plot.series:
                    out.append(f'  - {s.name}: ' + ', '.join('' if v is None else f'{v:g}' for v in s.values))
        except (IndexError, AttributeError, KeyError, ValueError, TypeError):
            out.append('  - (its numbers could not be read; read them off the picture)')
    if kind == MSO_SHAPE_TYPE.PICTURE:
        alt = sh._element.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr')
        d = clean(alt.get('descr', '')) if alt is not None else ''
        out.append(f'- [picture{": " + d if d else ""}]')
    return out


def outline(src, folder, pdf, wanted, say):
    pages_dir = folder / 'source' / 'pages'
    pages_dir.mkdir(parents=True, exist_ok=True)
    md = [f'# Outline of {src.name}', '']
    if src.suffix.lower() == '.pptx':
        from pptx import Presentation
        prs = Presentation(str(src))
        n = len(prs.slides)
        wanted = {k for k in wanted if 1 <= k <= n} if wanted else set(range(1, n + 1))
        facts = pptx_facts(prs, {})
        slides = {k: s for k, s in enumerate(prs.slides, 1) if k in wanted}
        hidden = {k for k in range(1, n + 1) if facts[k]['hidden']}
    else:
        import fitz
        doc = fitz.open(str(src))
        n = doc.page_count
        wanted = {k for k in wanted if 1 <= k <= n} if wanted else set(range(1, n + 1))
        slides, hidden, facts = {}, set(), {}
    pages = {}
    if pdf:
        import fitz
        import pptx2native as P
        tmp = tempfile.mkdtemp(prefix='importdeck-')
        try:
            pmap = page_map(n, hidden, fitz.open(str(pdf)).page_count, wanted)
            say(f'Rendering {len(pmap)} pages of {pdf.name}…')
            size = (prs.slide_width, prs.slide_height) if slides else pdf_size(fitz.open(str(pdf))[0])
            P.render_pdf(str(pdf), str(pages_dir), tmp, 'page', pmap, size)
            pages = {k: f'source/pages/page-{k:03d}.jpg' for k in pmap}
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
    md += [f'{len(wanted)} slides. Each has its title, its words in reading order (indented by bullet level), its tables and chart numbers, '
           'and its speaker notes. A picture of each slide is in source/pages/ when there was a PDF to render it from.', '']
    for k in sorted(wanted):
        if slides:
            s, f = slides[k], facts[k]
            body = []
            for sh in sorted(s.shapes, key=lambda c: (c.top or 0, c.left or 0)):
                if s.shapes.title is not None and sh.shape_id == s.shapes.title.shape_id:
                    continue
                body += shape_text(sh)
            title, notes = f['title'], f['notes']
            meta = [f'PowerPoint layout: {f["layout"]}'] + (['hidden in PowerPoint'] if f['hidden'] else []) + ([f'section: {f["section"]}'] if f['section'] else [])
        else:
            page = doc[k - 1]
            blocks = sorted((b for b in page.get_text('dict')['blocks'] if b.get('type') == 0), key=lambda b: (round(b['bbox'][1]), b['bbox'][0]))
            spans = [(sp['size'], clean(''.join(x['text'] for x in l['spans']))) for b in blocks for l in b['lines'] for sp in l['spans'][:1]]
            top = max(spans, default=(0, ''))
            title = top[1] if top[1] and len(top[1]) < 160 else ''
            body = ['- ' + t for _, t in spans if t and t != title]
            notes, meta = '', []
        md += [f'## Slide {k}: {title or "(no title)"}', '']
        if pages.get(k) or meta:
            md += [' · '.join(([f'Picture: {pages[k]}'] if pages.get(k) else []) + meta), '']
        md += (body or ['- (no words on the slide)']) + ['']
        if notes:
            md += ['Notes:', ''] + ['> ' + p for p in notes.splitlines() if p.strip()] + ['']
    out = folder / 'source' / 'outline.md'
    out.write_text('\n'.join(md), encoding='utf-8')
    return out, len(wanted), len(pages)


# ---------- main ----------

def mb(path):
    total = sum(p.stat().st_size for p in Path(path).rglob('*') if p.is_file())
    return f'{total / 1e6:.0f} MB' if total >= 1e6 else f'{total / 1e3:.0f} KB'


def main():
    ap = argparse.ArgumentParser(description='Bring a PowerPoint or PDF into a deck made with Import a deck.')
    ap.add_argument('deck', nargs='?')
    ap.add_argument('--check', action='store_true')
    ap.add_argument('--outline', action='store_true')
    ap.add_argument('--pdf')
    ap.add_argument('--no-pdf', action='store_true')
    ap.add_argument('--slides', default='')
    ap.add_argument('--no-wire', action='store_true')
    ap.add_argument('--force', action='store_true')
    a = ap.parse_args()
    say = lambda s: print(s, file=sys.stderr, flush=True)
    gone = missing()
    if a.check:
        print('Python packages: ' + ('all there' if not gone else 'missing ' + ', '.join(gone) + f'. Ask before installing them: {INSTALL}'))
        so = soffice()
        print(f'LibreOffice: {so}' if so else 'LibreOffice: not installed (a PDF of the deck has to come from PowerPoint, Keynote, or Google Slides)')
        sys.exit(1 if gone else 0)
    if not a.deck:
        ap.error('name the deck: its slug (./sb decks) or its folder')
    if gone:
        sys.exit(f'This needs {", ".join(gone)}. Ask the user before installing: {INSTALL}')
    deck, folder = find_deck(a.deck)
    src = source_of(deck, folder)
    if not src:
        sys.exit(f'No PowerPoint or PDF in {tilde(folder / "source")}. Put the file there, or start the import with ./sb new "<name>" --from <file>.')
    import pptx2native as P
    wanted = P.parse_slides(a.slides)
    if wanted and not (a.no_wire or a.outline):   # a few slides again: the deck file's other slides would be dropped
        sys.exit('With --slides, add --no-wire: the slides go into ct-slides.js and img/, and you add or update their LIBRARY entries.')
    deck_file = folder / deck.get('entry', 'deck.html')
    if not (a.no_wire or a.outline):
        imported_block(deck_file, a.force)
    pdf = pdf_for(src, a, say, required=not a.outline)

    if a.outline:
        out, n, pics = outline(src, folder, pdf, wanted, say)
        print(f'Wrote {tilde(out)}: {n} slides' + (', with a picture of each in source/pages/' if pics else ', without pictures (no PDF)') + '.')
        print('Rebuild each slide in the deck file with the layouts it carries, keeping its words and its notes. The skill\'s Importing a deck section has the steps.')
        return

    if src.suffix.lower() == '.pdf':
        data, pics, facts = convert_pdf(src, folder, wanted, say)
        stats = None
    else:
        conv, data, pics, facts = convert_pptx(src, folder, pdf, wanted, say)
        stats = conv.stats
    write_ct(folder, data, bool(wanted))
    if not a.no_wire:
        wire(deck_file, library(src, data, pics, facts), a.force, originals=bool(pics))

    # what came in
    editable = [k for k, v in data.items() if not v.get('error') and not v.get('bgpic') and v.get('shapes')]
    as_pic = {k: ('it could not be converted: ' + v['error'][:80]) if v.get('error') else 'its background is a picture that could not be read' if v.get('bgpic') else 'it has nothing to convert'
              for k, v in data.items() if k not in editable}
    untitled = [k for k in data if not facts[k]['title']]
    noted = [k for k in data if facts[k]['notes']]
    hidden = [k for k in data if facts[k]['hidden']]
    print(f'Brought in {len(data)} slides from source/{src.name}' + ('' if a.no_wire else f' into {deck.get("entry", "deck.html")}') + '.')
    print(f'  {len(editable)} editable: every word on them can be edited' + (' (a PDF: one text box per line)' if src.suffix.lower() == '.pdf' else ''))
    for k, why in sorted(as_pic.items()):
        print(f'  slide {k} as {"a picture of the original" if pics.get(k) else "an empty slide"}: {why}')
    if stats:
        print(f'  {stats.get("crops", 0)} charts, icons, or diagrams cut from the PDF as pictures' if pdf else '  no PDF, so any charts, SmartArt, and icons were left out')
    if src.suffix.lower() == '.pptx':
        print(f'  {len(noted)} with speaker notes' + (f'; {len(hidden)} hidden in PowerPoint, on the shelf (alt), outside the default version' if hidden else ''))
    if untitled:
        print(f'  no title found on slide{"s" * (len(untitled) > 1)} {", ".join(map(str, sorted(untitled)))}: named "Slide N" in the deck file; give each a title')
    print(f'  pictures: img/ ({mb(folder / "img")}); ' + ('each slide as it looked is img/imp-NNN.jpg, behind its "original" button' if pics else 'no pictures of the originals (no PDF)'))
    if not a.no_wire:
        print(f'Next: ./sb shot {deck.get("slug", folder.name)} <slide ids>, compare each with its original, then ./sb audit {deck.get("slug", folder.name)}.')


if __name__ == '__main__':
    main()
